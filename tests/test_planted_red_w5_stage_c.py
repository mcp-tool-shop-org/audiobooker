"""Planted-RED gates for wave-5 Stage C production contracts.

Written against the INTENDED contract. Production fixes land in sibling
worktrees (renderer-cache, parsers, cli-surface); until those merge, these
tests are expected RED. Do not xfail them.

Contracts:
  - F-ff5a0ed1  render, mutate opening text, sample, revert: skipped_cached==0
                and on-disk chapter WAV != sample bytes
  - F-c28baf2a  PDF outline prefix (pages before first bookmark) is kept
  - F-436150bd  DOCX footnote / endnote text is present in parsed chapters
  - F-1cb270de  unclosed aside + skip does not eat the rest of the chapter
  - F-d9513d2b  diagnose --json: json.loads(stdout) with no slice to first '{'
  - F-21e1ecd7  podcast feed has no empty enclosures
"""

from __future__ import annotations

import json
import sys
import types
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest

from audiobooker.parser.epub import html_to_text
from audiobooker.project import AudiobookProject
from audiobooker.renderer.cache_manifest import get_chapter_wav_path
from audiobooker.renderer.engine import render_project, render_sample
from audiobooker.renderer.output import AssemblyResult
from audiobooker.renderer.protocols import RunResult
from tests.fakes.fake_tts import FakeTTSEngine, write_silence_wav


# ---------------------------------------------------------------------------
# Optional extras
# ---------------------------------------------------------------------------

def _has(module: str) -> bool:
    try:
        __import__(module)
        return True
    except ImportError:
        return False


requires_pymupdf = pytest.mark.skipif(
    not _has("pymupdf"),
    reason="pymupdf not installed — install with: pip install pymupdf",
)
requires_python_docx = pytest.mark.skipif(
    not _has("docx"),
    reason="python-docx not installed — install with: pip install python-docx",
)


PROSE = (
    "The lamps along the quay had gone out one by one, and the harbour lay "
    "under a thin skin of ice that cracked whenever the swell came in "
    "beneath it. Mara counted the boats twice and got a different answer "
    "each time, which told her more about her hands than about the boats. "
)


def _prose(words: int, seed: str = "") -> str:
    base = (seed + " " + PROSE).split()
    out: list[str] = []
    while len(out) < words:
        out.extend(base)
    return " ".join(out[:words])


class RecordingAssembler:
    def __init__(self) -> None:
        self.calls: list[dict] = []

    def __call__(self, **kwargs) -> AssemblyResult:
        self.calls.append(kwargs)
        out = Path(kwargs["output_path"])
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_bytes(b"FAKE-AUDIO")
        return AssemblyResult(output_path=out, chapters_embedded=True)


class RecordingFfmpeg:
    def __init__(self) -> None:
        self.calls: list[list[str]] = []

    def run(self, args: list[str]) -> RunResult:
        self.calls.append(list(args))
        if args:
            last = args[-1]
            if last != "-" and not str(last).startswith("-"):
                p = Path(last)
                p.parent.mkdir(parents=True, exist_ok=True)
                p.write_bytes(b"FAKE-SAMPLE")
        return RunResult(returncode=0, stdout="", stderr="")


def _cast_narration_book(text: str, title: str = "StageC") -> AudiobookProject:
    project = AudiobookProject.from_string(text, title=title, author="Author")
    project.cast("narrator", "af_heart")
    project.compile()
    project.config.validate_voices_on_render = False
    return project


# ===========================================================================
# F-ff5a0ed1 — sample miss must not occupy the identity-keyed chapter WAV
# ===========================================================================

BOOK_TEXT = (
    "Chapter 1: The Harbor\n\n"
    "The tide came in slowly over the stones of the old quay. "
    "Morning found the boat empty and the lamp still turning."
)


class TestSampleMissDoesNotOccupyChapterWav:
    """renderer-cache F-ff5a0ed1 (CRITICAL).

    render_sample on a miss used to synthesize into chapter_NNNN.wav without
    rewriting the ChapterCacheEntry. Revert + render_project then reported
    Cached while the WAV was still the sample. Plant: render, mutate opening
    text, sample, revert, demand skipped_cached==0 and on-disk WAV != sample.

    Awaits sibling merge (renderer-cache).
    """

    def test_render_mutate_sample_revert_is_a_miss(self, tmp_path, monkeypatch):
        import audiobooker.renderer.output as output_mod

        monkeypatch.setattr(output_mod, "_ffmpeg_checked", True)
        monkeypatch.setattr(
            "audiobooker.renderer.ffmpeg_runner.RealFFmpegRunner",
            RecordingFfmpeg,
        )

        project = _cast_narration_book(BOOK_TEXT, title="SampleOccupy")
        cache_root = tmp_path / "cache"
        # Same duration so size_bytes collides; stock FakeTTS still varies PCM
        # by script (do not reintroduce a null oracle).
        engine = FakeTTSEngine(duration_per_call=1.0)
        first = render_project(
            project,
            tmp_path / "book.m4b",
            engine=engine,
            assembler=RecordingAssembler(),
            cache_root=cache_root,
        )
        wav = get_chapter_wav_path(cache_root, 0)
        assert wav.exists()
        original_text = project.chapters[0].utterances[0].text
        assert original_text

        project.chapters[0].utterances[0].text = (
            "A completely rewritten opening line about the Ashgate bell."
        )
        sample_engine = FakeTTSEngine(duration_per_call=1.0)
        render_sample(
            project,
            from_chapter=0,
            duration=30.0,
            output_path=tmp_path / "sample.m4a",
            engine=sample_engine,
            cache_root=cache_root,
        )
        sample_bytes = wav.read_bytes()

        project.chapters[0].utterances[0].text = original_text
        second_engine = FakeTTSEngine(duration_per_call=1.0)
        second = render_project(
            project,
            tmp_path / "book2.m4b",
            engine=second_engine,
            assembler=RecordingAssembler(),
            cache_root=cache_root,
        )
        summary = getattr(second, "render_summary", None)
        assert summary is not None, (
            "render_project returned a bare Path — skipped_cached is unreadable "
            f"(first={getattr(first, 'render_summary', None)!r})"
        )
        assert summary.skipped_cached == 0, (
            "revert after sample reported Cached — the sample occupied the "
            f"identity-keyed chapter WAV (skipped_cached={summary.skipped_cached} "
            f"rendered={summary.rendered}) (F-ff5a0ed1)"
        )
        assert wav.read_bytes() != sample_bytes, (
            "on-disk chapter WAV still equals the sample bytes after revert "
            "(F-ff5a0ed1)"
        )


# ===========================================================================
# F-c28baf2a — PDF outline prefix kept
# ===========================================================================

@requires_pymupdf
class TestPdfOutlinePrefixKept:
    """parsers F-c28baf2a. Awaits sibling merge (parsers)."""

    def test_dedication_before_first_bookmark_is_kept(self, tmp_path):
        import pymupdf
        from audiobooker.parser.pdf import parse_pdf

        path = tmp_path / "outlined.pdf"
        doc = pymupdf.open()
        dedication = (
            "DEDICATION_TOKEN this book is for the keepers of the Ashgate bell. "
            + _prose(80, "dedication")
        )
        page0 = doc.new_page()
        page0.insert_textbox(pymupdf.Rect(60, 60, 540, 740), dedication, fontsize=11)
        for n in (1, 2, 3):
            page = doc.new_page()
            page.insert_textbox(
                pymupdf.Rect(60, 60, 540, 740),
                _prose(80, f"chapter {n}"),
                fontsize=11,
            )
        # pymupdf TOC pages are 1-based; bookmarks at 2/3/4 leave page 0 unowned.
        doc.set_toc([
            [1, "Chapter 1", 2],
            [1, "Chapter 2", 3],
            [1, "Chapter 3", 4],
        ])
        doc.save(str(path))
        doc.close()

        _, chapters = parse_pdf(path, min_chapter_words=20)
        joined = "\n".join(c.raw_text for c in chapters)
        assert "DEDICATION_TOKEN" in joined, (
            "PDF outline dropped pages before the first bookmark: "
            f"titles={[c.title for c in chapters]!r} joined={joined[:240]!r} "
            "(F-c28baf2a)"
        )


# ===========================================================================
# F-436150bd — DOCX footnotes / endnotes walked
# ===========================================================================

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_OD_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _docx_with_notes(path: Path) -> Path:
    """Heading-1 document whose notes parts carry unique tokens."""
    import docx

    d = docx.Document()
    d.core_properties.title = "The Harbour Bell"
    d.add_heading("Chapter 1", level=1)
    d.add_paragraph(
        "The harbour bell rang over the ice and the keepers of Ashgate "
        "counted the boats twice before they believed the tide."
    )
    d.add_heading("Chapter 2", level=1)
    d.add_paragraph(
        "Morning found the quay empty again and the lamp unlit above the water."
    )
    d.save(str(path))

    footnotes_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:footnotes xmlns:w="{_W_NS}">'
        '<w:footnote w:type="separator" w:id="-1"><w:p/></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p/></w:footnote>'
        '<w:footnote w:id="1"><w:p><w:r>'
        "<w:t>FOOTNOTE_TOKEN the clapper faces north</w:t>"
        "</w:r></w:p></w:footnote>"
        "</w:footnotes>"
    )
    endnotes_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:endnotes xmlns:w="{_W_NS}">'
        '<w:endnote w:type="separator" w:id="-1"><w:p/></w:endnote>'
        '<w:endnote w:type="continuationSeparator" w:id="0"><w:p/></w:endnote>'
        '<w:endnote w:id="1"><w:p><w:r>'
        "<w:t>ENDNOTE_TOKEN the recast bell of 1842</w:t>"
        "</w:r></w:p></w:endnote>"
        "</w:endnotes>"
    )

    src = zipfile.ZipFile(path, "r")
    names = src.namelist()
    contents = {name: src.read(name) for name in names}
    src.close()

    ct = contents["[Content_Types].xml"].decode("utf-8")
    if "footnotes.xml" not in ct:
        ct = ct.replace(
            "</Types>",
            '<Override PartName="/word/footnotes.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument'
            '.wordprocessingml.footnotes+xml"/>'
            '<Override PartName="/word/endnotes.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument'
            '.wordprocessingml.endnotes+xml"/>'
            "</Types>",
        )
    contents["[Content_Types].xml"] = ct.encode("utf-8")

    rels_name = "word/_rels/document.xml.rels"
    rels = contents[rels_name].decode("utf-8")
    if "footnotes" not in rels:
        rels = rels.replace(
            "</Relationships>",
            f'<Relationship Id="rIdFootnotes" Type="{_OD_REL}/footnotes" '
            'Target="footnotes.xml"/>'
            f'<Relationship Id="rIdEndnotes" Type="{_OD_REL}/endnotes" '
            'Target="endnotes.xml"/>'
            "</Relationships>",
        )
    contents[rels_name] = rels.encode("utf-8")

    doc_xml = contents["word/document.xml"].decode("utf-8")
    marker = "</w:p>"
    note_refs = (
        '<w:r><w:footnoteReference w:id="1"/></w:r>'
        '<w:r><w:endnoteReference w:id="1"/></w:r>'
        "</w:p>"
    )
    # Attach the references to the first body paragraph close.
    if "<w:footnoteReference" not in doc_xml:
        doc_xml = doc_xml.replace(marker, note_refs, 1)
    contents["word/document.xml"] = doc_xml.encode("utf-8")
    contents["word/footnotes.xml"] = footnotes_xml.encode("utf-8")
    contents["word/endnotes.xml"] = endnotes_xml.encode("utf-8")

    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as out:
        for name, data in contents.items():
            out.writestr(name, data)
    return path


@requires_python_docx
class TestDocxNotesAreNarrated:
    """parsers F-436150bd. Awaits sibling merge (parsers)."""

    def test_footnote_and_endnote_tokens_are_present(self, tmp_path):
        from audiobooker.parser.docx import parse_docx

        path = _docx_with_notes(tmp_path / "notes.docx")
        _, chapters = parse_docx(path, min_chapter_words=10)
        joined = "\n".join(c.raw_text for c in chapters)
        assert "FOOTNOTE_TOKEN" in joined, (
            "DOCX footnote body was not walked: "
            f"titles={[c.title for c in chapters]!r} joined={joined[:240]!r} "
            "(F-436150bd)"
        )
        assert "ENDNOTE_TOKEN" in joined, (
            "DOCX endnote body was not walked: "
            f"titles={[c.title for c in chapters]!r} joined={joined[:240]!r} "
            "(F-436150bd)"
        )


# ===========================================================================
# F-1cb270de — unclosed aside + skip does not eat the rest of the chapter
# ===========================================================================

class TestUnclosedAsideSkipDoesNotEatChapter:
    """parsers F-1cb270de. Awaits sibling merge (parsers)."""

    def test_unclosed_footnote_aside_skip_keeps_after_text(self):
        html = (
            "<p>BEFORE the quay lamps.</p>"
            "<aside epub:type='footnote'>unclosed note about the clapper"
            "<p>AFTER UNIQUE the bell was recast in 1842.</p>"
        )
        text = html_to_text(html, footnote_behavior="skip")
        assert "BEFORE the quay lamps" in text, text
        assert "AFTER UNIQUE" in text, (
            "unclosed aside + skip deleted the rest of the chapter: "
            f"{text!r} (F-1cb270de)"
        )

    def test_unclosed_aside_inside_nav_does_not_eat_after_nav(self):
        html = (
            "<nav><aside epub:type='footnote'>nav note"
            "</nav>"
            "<p>AFTER UNIQUE the bell was recast in 1842.</p>"
        )
        text = html_to_text(html, footnote_behavior="skip")
        assert "AFTER UNIQUE" in text, (
            "unclosed aside inside <nav> plus skip deleted AFTER UNIQUE: "
            f"{text!r} (F-1cb270de)"
        )


# ===========================================================================
# F-d9513d2b — diagnose --json owns stdout
# ===========================================================================

class TestDiagnoseJsonOwnsStdout:
    """cli-surface F-d9513d2b. Awaits sibling merge (cli-surface)."""

    def test_diagnose_json_loads_stdout_without_slicing(
        self, monkeypatch, capsys
    ):
        from audiobooker.cli import main

        sys.modules.pop("fitz", None)

        real_import = __import__

        def noisy_import(name, globals=None, locals=None, fromlist=(), level=0):
            if name == "fitz":
                print(
                    "warning: The `fitz` API is deprecated. "
                    "Use `import pymupdf` instead."
                )
                fake = types.ModuleType("fitz")
                sys.modules["fitz"] = fake
                return fake
            return real_import(name, globals, locals, fromlist, level)

        monkeypatch.setattr("builtins.__import__", noisy_import)

        code = main(["diagnose", "--json"])
        captured = capsys.readouterr()
        # Must parse WITHOUT slicing to the first '{'.
        payload = json.loads(captured.out)
        assert isinstance(payload, dict)
        assert "ready" in payload
        stripped = captured.out.strip()
        assert stripped.startswith("{"), captured.out[:80]
        assert stripped.endswith("}")
        assert (code == 0) is bool(payload["ready"])


# ===========================================================================
# F-21e1ecd7 — podcast feed has no empty enclosures
# ===========================================================================

class TestPodcastNoEmptyEnclosures:
    """cli-surface / project-core F-21e1ecd7. Awaits sibling merge."""

    def test_partial_render_does_not_emit_empty_enclosure(
        self, tmp_path, capsys, monkeypatch
    ):
        from audiobooker.cli import main

        monkeypatch.chdir(tmp_path)
        text = (
            "Chapter 1: One\n\n"
            "The harbour lay under a thin skin of ice that cracked.\n\n"
            "Chapter 2: Two\n\n"
            "Morning found the boat empty and the lamp still turning."
        )
        project = _cast_narration_book(text, title="Feed")
        wav = tmp_path / "ch01.mp3"
        write_silence_wav(wav, duration_s=0.4, marker=0xC0FFEE)
        project.chapters[0].audio_path = wav
        project.chapters[0].duration_seconds = 0.4
        project.chapters[1].audio_path = None
        path = tmp_path / "feed.audiobooker"
        project.save(path)

        out_xml = tmp_path / "podcast.xml"
        code = main([
            "podcast", "-p", str(path),
            "--no-render",
            "--base-url", "https://x.test/feed/",
            "-o", str(out_xml),
        ])
        assert code == 0, capsys.readouterr()
        xml = out_xml.read_text(encoding="utf-8")
        root = ET.fromstring(xml)
        items = root.find("channel").findall("item")
        enclosures = [
            item.find("enclosure") for item in items if item.find("enclosure") is not None
        ]
        assert enclosures, xml
        for enc in enclosures:
            url = enc.attrib.get("url", "")
            filename = url.rsplit("/", 1)[-1]
            assert filename, (
                f"podcast feed emitted an empty enclosure url={url!r} (F-21e1ecd7)"
            )
            assert url.rstrip("/") != "https://x.test/feed"
        assert len(items) == len(enclosures)
        assert len(items) == 1, (
            "partial --no-render feed kept gap chapters as empty enclosures: "
            f"{len(items)} items (F-21e1ecd7)"
        )
